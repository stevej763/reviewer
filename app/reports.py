"""Report generation and exports (Markdown, HTML, Word, CSV)."""

import csv
import html
import io
import re
from datetime import date

import markdown as md
from docx import Document
from docx.shared import Pt

from .prompts import REPORT_STYLES, REPORT_SYSTEM


def _fmt_list(label: str, items: list[str]) -> str:
    return f"{label}: {'; '.join(items)}\n" if items else ""


def entries_digest(entries: list[dict]) -> str:
    """Compact, chronological text of summaries to hand to the model."""
    lines = []
    for entry in sorted(entries, key=lambda e: (e["date"], e["id"])):
        s = entry.get("summary")
        weekday = date.fromisoformat(entry["date"]).strftime("%a")
        lines.append(f"## {weekday} {entry['date']}")
        if not s:
            lines.append(entry["text"].strip())
            lines.append("")
            continue
        if s.get("headline"):
            lines.append(s["headline"])
        for h in s.get("highlights", []):
            impact = f" (impact: {h['impact']})" if h.get("impact") else ""
            lines.append(f"- [{h['category']}] {h['text']}{impact}")
        block = (
            _fmt_list("Wins", s.get("wins", []))
            + _fmt_list("Projects", s.get("projects", []))
            + _fmt_list("People", s.get("people", []))
            + _fmt_list("Challenges", s.get("challenges", []))
            + _fmt_list("Follow-ups", s.get("follow_ups", []))
        )
        if block:
            lines.append(block.rstrip())
        lines.append("")
    return "\n".join(lines)


def report_prompt(entries: list[dict], start: str, end: str, style: str, notes: str = "") -> tuple[str, str]:
    spec = REPORT_STYLES[style]
    prompt = (
        f"Period: {start} to {end} ({len(entries)} log entries).\n\n"
        f"Report type: {spec['label']}\n{spec['instructions']}\n"
    )
    if notes.strip():
        prompt += f"\nExtra guidance from me:\n{notes.strip()}\n"
    prompt += f"\n<work_log>\n{entries_digest(entries)}\n</work_log>"
    return REPORT_SYSTEM, prompt


def title_from_markdown(content: str, fallback: str) -> str:
    match = re.search(r"^#\s+(.+)$", content, re.MULTILINE)
    return match.group(1).strip() if match else fallback


def render_markdown(content: str) -> str:
    """Markdown to HTML with raw HTML disabled: escaping '&' and '<' means no tags
    can come from the model's output, while '>' blockquotes still work."""
    safe = content.replace("&", "&amp;").replace("<", "&lt;")
    return md.markdown(safe, extensions=["tables", "fenced_code", "sane_lists"])


def to_html(title: str, content: str) -> str:
    body = render_markdown(content)
    return f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>{html.escape(title)}</title>
<style>
body{{font:16px/1.6 -apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,sans-serif;max-width:760px;margin:40px auto;padding:0 20px;color:#1d2433}}
h1{{font-size:2rem;line-height:1.2}} h2{{margin-top:2rem;border-bottom:1px solid #e3e6ec;padding-bottom:.3rem}}
li{{margin:.25rem 0}} blockquote{{border-left:3px solid #c9cfda;margin:0;padding-left:1rem;color:#4a5468}}
</style></head><body>{body}</body></html>"""


def _add_runs(paragraph, text: str):
    # Minimal inline formatting: **bold** and *italic*.
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        elif part:
            paragraph.add_run(part)


def to_docx(content: str) -> bytes:
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)
    for line in content.splitlines():
        stripped = line.strip()
        if not stripped or stripped in {"---", "***"}:
            continue
        heading = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        bullet = re.match(r"^[-*+]\s+(.*)$", stripped)
        numbered = re.match(r"^\d+[.)]\s+(.*)$", stripped)
        indent = len(line) - len(line.lstrip())
        if heading:
            doc.add_heading(heading.group(2).replace("**", ""), level=min(len(heading.group(1)) - 1, 3))
        elif bullet:
            style = "List Bullet 2" if indent >= 2 else "List Bullet"
            _add_runs(doc.add_paragraph(style=style), bullet.group(1))
        elif numbered:
            _add_runs(doc.add_paragraph(style="List Number"), numbered.group(1))
        elif stripped.startswith(">"):
            _add_runs(doc.add_paragraph(style="Quote"), stripped.lstrip("> "))
        else:
            _add_runs(doc.add_paragraph(), stripped)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def highlights_csv(entries: list[dict]) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["date", "category", "highlight", "impact", "projects", "people", "tags"])
    for entry in sorted(entries, key=lambda e: (e["date"], e["id"])):
        s = entry.get("summary") or {}
        common = ["; ".join(s.get("projects", [])), "; ".join(s.get("people", [])), "; ".join(s.get("tags", []))]
        for h in s.get("highlights", []):
            writer.writerow([entry["date"], h["category"], h["text"], h.get("impact", ""), *common])
    return buf.getvalue()
